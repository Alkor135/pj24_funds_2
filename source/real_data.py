"""Чтение исходных XLSX и причинный календарь исследования Real.

Логика: соединяет «Прогноз» и «Диагностику» по уникальному ключу,
проверяет вероятности и числа, отделяет признаки до открытия от фактов.
Общий календарь обеих серий разбивается 60/20/20 до фильтрации результатов.
Основные расходы — 3 б.п. на круг. Модуль пакета source вызывается корневым
скриптом research_real.py.

Примеры из папки проекта:
    python research_real.py
    python research_real.py --cost-bps 3 --min-history 200
"""

from __future__ import annotations

import hashlib
import json
import math
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

VERSION = 'real-research-1.0'
STATES = ('Бычий', 'Медвежий', 'Боковой')
MOVEMENTS = ('Рост', 'Падение', 'Флэт')
FEATURE_NAMES = {'edge_direction': 'Перевес направления', 'p_flat': 'Вероятность Флэт',
                 'edge_regime': 'Перевес режима', 'p_side': 'Вероятность бокового режима',
                 'sigma_oc': 'Волатильность OPEN → CLOSE', 'g': 'Предыдущее движение режима',
                 'z': 'Нормированное движение режима'}
PERIOD_NAMES = {'train': 'Обучение', 'validation': 'Выбор', 'test': 'Тест'}


@dataclass(frozen=True)
class ResearchConfig:
    """Параметры исследования: расходы, минимумы сделок и блочная проверка."""
    cost_bps: float = 3.0
    cost_scenarios_bps: tuple[float, ...] = (0, 2, 3, 5, 10)
    min_history: int = 200
    min_train_trades: int = 60
    min_validation_trades: int = 30
    min_test_trades: int = 30
    bootstrap_repeats: int = 2000
    seed: int = 42

    def __post_init__(self):
        """Проверяет параметры экземпляра; при недопустимом значении вызывает ValueError."""
        costs = (self.cost_bps, *self.cost_scenarios_bps)
        if not costs or any(not math.isfinite(v) or v < 0 for v in costs):
            raise ValueError('Расходы должны быть конечными неотрицательными числами.')
        if any(not isinstance(v, int) or isinstance(v, bool) or v < 1 for v in
               (self.min_history, self.min_train_trades, self.min_validation_trades,
                self.min_test_trades, self.bootstrap_repeats)):
            raise ValueError('Минимальные размеры выборок и число повторов должны быть положительными целыми.')
        if not isinstance(self.seed,int) or isinstance(self.seed,bool) or self.seed<0:
            raise ValueError('Seed должен быть неотрицательным целым числом.')

    @property
    def scenarios(self) -> tuple[float, ...]:
        """Возвращает уникальные расходы, включая основные и стрессовые, по возрастанию."""
        return tuple(sorted(set((*self.cost_scenarios_bps, self.cost_bps, self.stress_cost))))

    @property
    def stress_cost(self) -> float:
        """Возвращает стрессовые расходы на круг в б.п., не меньше 10 и удвоенных основных."""
        return max(10.0, 2 * self.cost_bps)


@dataclass
class Record:
    """Дневная строка: признаки и вероятности отделены от итоговых фактов и статусов."""
    day: date
    secid: str | None = None
    expiry: date | None = None
    features: dict = field(default_factory=dict)
    p_d: tuple | None = None
    p_r: tuple | None = None
    state: str | None = None
    n_d: int = 0
    n_r: int = 0
    source_d: str | None = None
    source_r: str | None = None
    real_points: float | None = None
    r_oc: float | None = None
    actual_d: int | None = None
    actual_r: int | None = None
    matrix_d: tuple = ((0, 0, 0),) * 3
    matrix_r: tuple = ((0, 0, 0),) * 3
    short_window: bool = False
    statuses: tuple[str, ...] = ()


@dataclass
class Dataset:
    """История одной серии с метаданными, исходным путём и проверенным SHA-256."""
    series: str
    records: list[Record]
    metadata: dict = field(default_factory=dict)
    sha256: str = ''
    path: Path | None = None


def file_hash(path: Path) -> str:
    """Возвращает SHA-256 содержимого файла path, читая его блоками по 1 МиБ."""
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1024*1024), b''):
            digest.update(block)
    return digest.hexdigest()


def finite_number(value, name: str, optional=True) -> float | None:
    """Проверяет число value поля name; возвращает float либо разрешённый None."""
    if value is None and optional:
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError(f'{name}: требуется конечное число, получено {value!r}.')
    return float(value)


def probabilities(values, name: str) -> tuple | None:
    """Возвращает проверенную тройку values поля name либо None для полностью пустой."""
    if all(v is None for v in values):
        return None
    result = tuple(finite_number(v, name, False) for v in values)
    if any(v < 0 or v > 1 for v in result) or abs(sum(result)-1) > 1e-10:
        raise ValueError(f'{name}: вероятности должны быть в [0,1] и суммироваться в 1.')
    return result


def as_date(value, name: str, optional=False) -> date | None:
    """Возвращает дату value поля name; пустую допускает только при optional=True."""
    if value is None and optional:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            pass
    raise ValueError(f'{name}: требуется календарная дата.')


def count_value(value, name: str) -> int:
    """Возвращает целый неотрицательный счётчик value поля name или вызывает ошибку."""
    number = finite_number(value, name, False)
    if number < 0 or number != int(number):
        raise ValueError(f'{name}: требуется неотрицательный целый счётчик.')
    return int(number)


def count_matrix(value, name: str) -> tuple:
    """Разбирает JSON value поля name в матрицу 3×3 целых неотрицательных счётчиков."""
    if value is None:
        return ((0, 0, 0),)*3
    try:
        matrix = json.loads(value)
        if len(matrix) != 3 or any(len(row) != 3 for row in matrix):
            raise ValueError
        return tuple(tuple(count_value(v, name) for v in row) for row in matrix)
    except (TypeError, json.JSONDecodeError, ValueError) as error:
        raise ValueError(f'{name}: некорректная матрица счётчиков.') from error


def keyed_rows(sheet, required: set[str]) -> dict:
    """Возвращает строки sheet по ключу дата/SECID/экспирация, проверяя required и дубликаты."""
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator, ())
    if len(headers) != len(set(headers)) or not required.issubset(headers):
        raise ValueError(f'Лист {sheet.title}: неверные или повторяющиеся заголовки.')
    result = {}
    for values in iterator:
        if all(v is None for v in values):
            continue
        row = dict(zip(headers, values))
        key = (as_date(row['Дата'], 'Дата'), row['SECID'], as_date(row['LSTTRADE'], 'LSTTRADE', True))
        if key in result:
            raise ValueError(f'Лист {sheet.title}: повтор ключа {key}.')
        result[key] = row
    return result


def class_index(value, labels: tuple, name: str) -> int | None:
    """Возвращает индекс метки value в labels, разрешая пустое поле name."""
    if value is None:
        return None
    if value not in labels:
        raise ValueError(f'{name}: неизвестная метка {value!r}.')
    return labels.index(value)


def load_dataset(path: Path, series: str) -> Dataset:
    """Читает и проверяет книгу path серии series; возвращает Dataset без записи источника."""
    from .funds_xlsx import HEADERS, DETAILS
    path = Path(path).resolve()
    before = file_hash(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if not {'Прогноз', 'Диагностика', 'Параметры'}.issubset(book.sheetnames):
            raise ValueError('Требуются листы Прогноз, Диагностика и Параметры.')
        main = keyed_rows(book['Прогноз'], set(HEADERS))
        diagnostics = keyed_rows(book['Диагностика'],
            {'Дата', 'SECID', 'LSTTRADE', 'p_up', 'p_down', 'p_flat', 'p_bull', 'p_bear', 'p_side',
             *(label for key, label in DETAILS)})
        if not main or main.keys() != diagnostics.keys():
            raise ValueError('Состав ключей Прогноза и Диагностики должен совпадать и быть непустым.')
        metadata = dict(book['Параметры'].iter_rows(min_row=2, max_col=2, values_only=True))
        if metadata.get('Серия') != series:
            raise ValueError(f'Книга не соответствует серии {series}.')
        length = count_value(metadata.get('L'), 'L')
        alpha = finite_number(metadata.get('alpha'), 'alpha', False)
        regime_window = count_value(metadata.get('W_max',60),'W_max')
        movement_window = count_value(metadata.get('V_max',60),'V_max')
        if length < 1 or alpha < 0:
            raise ValueError('Некорректные параметры L или alpha.')
        if min(regime_window,movement_window)<2:
            raise ValueError('Максимальные окна должны содержать хотя бы две доходности.')
        result = []
        for key in sorted(main, key=lambda k:k[0]):
            a, b = main[key], diagnostics[key]
            probabilities([a[n] for n in HEADERS[3:6]], 'Отображённое движение')
            probabilities([a[n] for n in HEADERS[6:9]], 'Отображённый режим')
            p_d = probabilities([b[n] for n in ('p_up','p_down','p_flat')], 'Движение')
            p_r = probabilities([b[n] for n in ('p_bull','p_bear','p_side')], 'Режим')
            sigma = finite_number(b['sigma_CC предыдущего режима'], 'sigma_CC')
            g = finite_number(b['g предыдущего режима'], 'g')
            oc = finite_number(b['sigma_OC'], 'sigma_OC')
            if any(v is not None and v < 0 for v in (sigma, oc)):
                raise ValueError('Волатильность не может быть отрицательной.')
            y = finite_number(b['Факт ln(CLOSE/OPEN)'], 'Факт доходности')
            r = math.expm1(y) if y is not None else None
            if r is not None and (not math.isfinite(r) or r <= -1):
                raise ValueError('Некорректный фактический результат OPEN → CLOSE.')
            real = finite_number(a['Real'], 'Real')
            if real is not None and r is not None and ((real>0 and r<0) or (real<0 and r>0) or (real==0 and r!=0)):
                raise ValueError('Знак Real не соответствует фактической доходности.')
            state = b['Наблюдаемый режим']
            class_index(state, STATES, 'Наблюдаемый режим')
            actual_d = class_index(b['Факт движения'], MOVEMENTS, 'Факт движения')
            actual_r = class_index(b['Факт режима'], STATES, 'Факт режима')
            epsilon = finite_number(b['epsilon движения'], 'epsilon')
            if y is not None and epsilon is not None and actual_d is not None:
                expected = 0 if y > epsilon else 1 if y < -epsilon else 2
                if actual_d != expected:
                    raise ValueError('Метка движения не соответствует сохранённой границе.')
            features = {'edge_direction': p_d[0]-p_d[1] if p_d else None,
                        'p_flat': p_d[2] if p_d else None,
                        'edge_regime': p_r[0]-p_r[1] if p_r else None,
                        'p_side': p_r[2] if p_r else None, 'sigma_oc': oc, 'g': g,
                        'z': g/(sigma*math.sqrt(length)) if g is not None and sigma and sigma>0 else None}
            windows = [count_value(b[n],n) if b[n] is not None else None
                       for n in ('W_eff предыдущего режима','W_eff целевого режима','V_eff')]
            maximums = (regime_window,regime_window,movement_window)
            if any(v is not None and v>maximum for v,maximum in zip(windows,maximums)):
                raise ValueError('Размер окна превышает сохранённый максимум модели.')
            n_d, n_r = count_value(b['n_D'],'n_D'), count_value(b['n_R'],'n_R')
            matrix_d = count_matrix(b[DETAILS[23][1]],'Матрица D')
            matrix_r = count_matrix(b[DETAILS[24][1]],'Матрица P1')
            for matrix, count, probability, total_name in (
                (matrix_d,n_d,p_d,'Всего обучающих исходов D'),
                (matrix_r,n_r,p_r,'Всего обучающих исходов P1')):
                total = sum(map(sum,matrix))
                saved_total = count_value(b[total_name],total_name)
                counts = matrix[STATES.index(state)] if state in STATES else (0,0,0)
                if not sum(counts):
                    counts = tuple(sum(row[c] for row in matrix) for c in range(3))
                if total!=saved_total or sum(counts)!=count:
                    raise ValueError('Матрица счётчиков не соответствует сохранённым размерам истории.')
                if probability is not None:
                    if count==0 or any(abs(p-(v+alpha)/(count+3*alpha))>1e-10
                                       for p,v in zip(probability,counts)):
                        raise ValueError('Прогноз не соответствует сохранённой матрице и сглаживанию.')
            result.append(Record(key[0], key[1], key[2], features, p_d, p_r, state,
                                 n_d, n_r,
                                 b['Источник движения'], b['Источник режима'], real, r, actual_d, actual_r,
                                 matrix_d, matrix_r,
                                 any(v is not None and v<maximum for v,maximum in zip(windows,maximums)),
                                 tuple(s.strip() for s in str(b['Статусы'] or '').split(';') if s.strip())))
    finally:
        book.close()
    if file_hash(path) != before:
        raise ValueError('Исходная книга изменилась во время чтения.')
    if len({r.day for r in result}) != len(result):
        raise ValueError('В серии должен быть один основной контракт на дату.')
    return Dataset(series, result, metadata, before, path)


def split_calendar(datasets: list[Dataset]) -> dict[str, list[date]]:
    """Возвращает общее хронологическое разбиение 60/20/20 записей datasets до фильтрации."""
    days = sorted({row.day for dataset in datasets for row in dataset.records})
    if len(days) < 5:
        raise ValueError('Для трёх временных участков требуется хотя бы пять дат.')
    a, b = math.floor(0.6*len(days)), math.floor(0.8*len(days))
    return {'train':days[:a], 'validation':days[a:b], 'test':days[b:]}


def eligible(record: Record, config: ResearchConfig, require_regime=False) -> bool:
    """Возвращает пригодность признаков record по config до открытия; факты не проверяет."""
    direction = record.p_d is not None and record.source_d=='CONDITIONAL_STATE' and record.n_d>=config.min_history
    regime = record.p_r is not None and record.source_r=='CONDITIONAL_STATE' and record.n_r>=config.min_history and record.state in STATES
    return direction and (not require_regime or regime)
