"""Сохранение дневных вероятностей в XLSX с диагностикой и параметрами.

Логика: лист «Прогноз» имеет десять колонок, включая Флэт и Real=CLOSE−OPEN
в пунктах, типизированные даты и числовые проценты. Отображённые тройки
округляются до 0.1 процентного
пункта методом наибольших остатков, сохраняя сумму 100.0%. Лист «Диагностика»
содержит исходные доли, пороги, счётчики, статусы и факты. Параметры и источник
фиксируются отдельно. По умолчанию существующий файл защищён от перезаписи.

Этот модуль пакета source используется корневыми точками запуска;
зависимость: openpyxl>=3.1.5,<4.
Примеры из папки проекта:
    python -m pip install -r requirements.txt
    python forecast_rts.py --output outputs/RTS_probabilities.xlsx
    python forecast_mix.py --output outputs/MIX_probabilities.xlsx --overwrite
"""

from __future__ import annotations

import json
import math
import os
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Sequence

from .funds_model import Forecast, STATE_NAMES, MOVEMENT_NAMES


HEADERS = ("Дата", "SECID", "LSTTRADE", "Рост", "Падение",
           "Флэт", "Бычий", "Медвежий", "Боковой", "Real")

DETAILS = (
    ("signal_day", "Предыдущая дата"), ("signal_secid", "SECID истории"),
    ("signal_state", "Наблюдаемый режим"), ("movement_source", "Источник движения"),
    ("regime_source", "Источник режима"), ("n_D", "n_D"), ("n_R", "n_R"),
    ("train_D", "Всего обучающих исходов D"), ("train_R", "Всего обучающих исходов P1"),
    ("signal_g", "g предыдущего режима"), ("signal_sigma", "sigma_CC предыдущего режима"),
    ("signal_window", "W_eff предыдущего режима"), ("target_sigma", "sigma_CC целевого режима"),
    ("target_window", "W_eff целевого режима"), ("target_boundary", "b целевого режима"),
    ("target_anchor", "Опорное CLOSE для цели"), ("lower_close", "Нижняя граница CLOSE"),
    ("upper_close", "Верхняя граница CLOSE"), ("oc_sigma", "sigma_OC"),
    ("oc_window", "V_eff"), ("epsilon", "epsilon движения"),
    ("lower_movement", "Нижняя граница движения, %"), ("upper_movement", "Верхняя граница движения, %"),
    ("matrix_D", "Счётчики D: строки Бычий/Медвежий/Боковой, столбцы Рост/Падение/Флэт"),
    ("matrix_R", "Счётчики P1: строки и столбцы Бычий/Медвежий/Боковой"),
    ("actual_y", "Факт ln(CLOSE/OPEN)"), ("actual_movement", "Факт движения"),
    ("actual_regime", "Факт режима"), ("filled_target", "Перенесённый целевой бар"),
    ("statuses", "Статусы"),
)


def rounded_probabilities(values: Sequence[float]) -> tuple[float, float, float]:
    """Округляет три доли values с суммой 1 до десятых процента.

    Возвращает доли с суммой 1; равные остатки разрешаются в исходном порядке.
    """
    if len(values) != 3 or any(not math.isfinite(p) or p < 0 for p in values) or abs(sum(values) - 1) > 1e-12:
        raise ValueError("Для округления нужна тройка неотрицательных долей с суммой 1.")
    scaled = [p * 1000 for p in values]
    units = [math.floor(p) for p in scaled]
    remainder = 1000 - sum(units)
    order = sorted(range(3), key=lambda i: (-(scaled[i] - units[i]), i))
    for i in order[:remainder]:
        units[i] += 1
    return tuple(value / 1000 for value in units)


def format_sheet(sheet, widths: Sequence[float], header_height: float = 46):
    """Оформляет sheet с ширинами widths и высотой заголовка header_height.

    Закрепляет первую строку, добавляет фильтр; результат отсутствует.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    sheet.freeze_panes = "D2" if sheet.title == "Прогноз" else "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = header_height
    sheet.print_title_rows = "1:1"
    for cell in sheet[1]:
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24476B")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="243746")
            cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(cell.value, str) else "right")
            if isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, float):
                cell.number_format = "0.000000"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def diagnostic_value(key: str, value):
    """Преобразует value поля key в подпись или JSON для листа диагностики."""
    if value is None:
        return None
    if key in ("signal_state", "actual_regime"):
        return STATE_NAMES[value]
    if key == "actual_movement":
        return MOVEMENT_NAMES[value]
    if key.startswith("matrix_"):
        return json.dumps(value, ensure_ascii=False)
    return value


def write_workbook(path: Path, rows: Sequence[Forecast], metadata: dict,
                   overwrite: bool = False) -> None:
    """Сохраняет rows и metadata в XLSX по path с флагом overwrite.

    Создаёт папку результата, готовит XLSX во временном файле в той же папке;
    возвращаемого значения нет. Без overwrite существующий файл даёт ошибку.
    Флэт хранится как вероятность, Real из rows.details — как число пунктов
    либо пустое значение для отсутствующего или синтетического факта.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, PatternFill
        from openpyxl.formatting.rule import ColorScaleRule
    except ImportError as error:
        raise ImportError("Установите зависимость: python -m pip install -r requirements.txt") from error
    path = Path(path).resolve()
    if path.exists() and not overwrite:
        raise FileExistsError(f"Файл существует: {path}. Перезапись требует --overwrite.")
    if len(rows) > 1_048_575:
        raise ValueError("Число строк превышает вместимость листа Excel.")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    main = workbook.active
    main.title = "Прогноз"
    main.append(HEADERS)
    audit = workbook.create_sheet("Диагностика")
    audit.append(("Дата", "SECID", "LSTTRADE", "p_up", "p_down", "p_flat",
                  "p_bull", "p_bear", "p_side") + tuple(label for key, label in DETAILS))
    for forecast in rows:
        identifiers = (forecast.day, forecast.secid, forecast.expiry)
        rounded_d = rounded_probabilities(forecast.movement) if forecast.movement is not None else (None,) * 3
        rounded_r = rounded_probabilities(forecast.regime) if forecast.regime is not None else (None,) * 3
        main.append(identifiers + rounded_d + rounded_r + (forecast.details.get("real_points"),))
        audit.append(identifiers + (forecast.movement or (None,) * 3)
                     + (forecast.regime or (None,) * 3)
                     + tuple(diagnostic_value(key, forecast.details.get(key)) for key, label in DETAILS))

    format_sheet(main, (14, 19, 14, 13, 13, 13, 13, 13, 13, 17))
    for row in main.iter_rows(min_row=2, min_col=10, max_col=10):
        row[0].number_format = "#,##0.00"
    for row in main.iter_rows(min_row=2, min_col=4, max_col=9):
        for cell in row:
            cell.number_format = "0.0%"
            if cell.value is None:
                cell.fill = PatternFill("solid", fgColor="E9EDF1")
    if rows:
        for area in (f"D2:F{main.max_row}", f"G2:I{main.max_row}"):
            main.conditional_formatting.add(area, ColorScaleRule(start_type="num", start_value=0,
                start_color="F5F8FC", end_type="num", end_value=1, end_color="7AAED3"))
    main.sheet_view.zoomScale = 90
    audit_widths = [14, 19, 14] + [16] * 6 + [max(16, min(32, len(label) * 0.8)) for key, label in DETAILS]
    audit_widths[-1] = 72
    audit_widths[-7] = 52
    audit_widths[-6] = 52
    audit_widths[12] = 34
    audit_widths[13] = 34
    format_sheet(audit, audit_widths, 76)
    for row in audit.iter_rows(min_row=2):
        row[-1].alignment = Alignment(wrap_text=True, vertical="center")
        audit.row_dimensions[row[0].row].height = max(32, 16 * math.ceil(len(row[-1].value or "") / 85))
        for cell in (row[12], row[13]):
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        for column in range(4, 10):
            row[column - 1].number_format = "0.000000%"
        for offset, (key, label) in enumerate(DETAILS, 10):
            if key in ("lower_movement", "upper_movement"):
                row[offset - 1].number_format = "0.0000%"

    settings = workbook.create_sheet("Параметры")
    settings.append(("Параметр", "Значение"))
    for key, value in metadata.items():
        settings.append((key, value))
    format_sheet(settings, (34, 125), 30)
    settings.auto_filter.ref = None
    for row in settings.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        settings.row_dimensions[row[0].row].height = max(24, 16 * math.ceil(len(str(row[1].value or "")) / 105))
    workbook.properties.creator = "pj24_funds_2"
    workbook.properties.title = "Дневные вероятности движения и режима"
    workbook.properties.description = "Полный дневной бар Day. Доли модели и причинная историческая реконструкция."
    handle, temporary = tempfile.mkstemp(prefix=".probabilities-", suffix=".xlsx", dir=path.parent)
    os.close(handle)
    temporary_path = Path(temporary)
    try:
        workbook.save(temporary_path)
        if overwrite:
            temporary_path.replace(path)
        else:
            # Режим xb защищает также от файла, появившегося во время расчёта.
            created = False
            try:
                with path.open("xb") as destination:
                    created = True
                    with temporary_path.open("rb") as source:
                        shutil.copyfileobj(source, destination)
            except OSError:
                if created:
                    path.unlink(missing_ok=True)
                raise
    finally:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
