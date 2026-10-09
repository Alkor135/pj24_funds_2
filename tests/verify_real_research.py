r"""Независимая сверка готового исследования Real с исходными XLSX.

Логика: стандартная библиотека и openpyxl повторно проверяют исходные хеши,
календарь 60/20/20, условия сигналов, расходы на круг, агрегаты и Brier.
Модули исследуемой реализации не импортируются. Проверка только читает файлы.

Примеры из папки проекта:
    .venv\Scripts\python.exe tests/verify_real_research.py
    python tests/verify_real_research.py --summary outputs/real_research/summary.json
"""

import argparse
import csv
import hashlib
import json
import math
import statistics
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    """Возвращает SHA-256 файла path без изменения его содержимого."""
    digest=hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda:stream.read(1024*1024),b''):
            digest.update(block)
    return digest.hexdigest()


def source_rows(path: Path) -> list[dict]:
    """Независимо читает исходную книгу path и возвращает факты, прогнозы и матрицы по датам."""
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        parts=[]
        for name in ('Прогноз','Диагностика'):
            iterator=book[name].iter_rows(values_only=True)
            headers=next(iterator)
            parts.append({str(row[0].date()):dict(zip(headers,row)) for row in iterator})
        assert parts[0].keys()==parts[1].keys(),'Ключи листов различаются'
        result=[]
        for day,a in parts[0].items():
            b=parts[1][day]
            p=tuple(b[k] for k in ('p_up','p_down','p_flat'))
            regime=tuple(b[k] for k in ('p_bull','p_bear','p_side'))
            y=b['Факт ln(CLOSE/OPEN)']
            result.append({'date':day,'p_d':p,'p_r':regime,
                           'edge':p[0]-p[1] if all(v is not None for v in p) else None,
                           'flat':p[2],'state':b['Наблюдаемый режим'],'nd':b['n_D'],'nr':b['n_R'],
                           'source_d':b['Источник движения'],'source_r':b['Источник режима'],
                           'r':math.expm1(y) if y is not None else None,'real':a['Real'],
                           'actual_d':b['Факт движения'],'actual_r':b['Факт режима'],
                           'matrix_d':json.loads(b['Счётчики D: строки Бычий/Медвежий/Боковой, столбцы Рост/Падение/Флэт'])
                               if b['Счётчики D: строки Бычий/Медвежий/Боковой, столбцы Рост/Падение/Флэт'] else [[0]*3 for _ in range(3)],
                           'matrix_r':json.loads(b['Счётчики P1: строки и столбцы Бычий/Медвежий/Боковой'])
                               if b['Счётчики P1: строки и столбцы Бычий/Медвежий/Боковой'] else [[0]*3 for _ in range(3)]})
        return sorted(result,key=lambda r:r['date'])
    finally:
        book.close()


def available(row: dict, config: dict, regime=False) -> bool:
    """Возвращает причинную пригодность исходной row по config, дополнительно режима при regime=True."""
    direction=all(v is not None for v in row['p_d']) and row['source_d']=='CONDITIONAL_STATE' and row['nd']>=config['min_history']
    state=all(v is not None for v in row['p_r']) and row['source_r']=='CONDITIONAL_STATE' and row['nr']>=config['min_history']
    return direction and (not regime or state)


def signal(row: dict, rule: dict | None, config: dict) -> int:
    """Возвращает независимый сигнал исходной row по rule/config, без обращения к её факту."""
    if rule is None or not available(row,config,rule['state'] is not None):
        return 0
    if rule['relation']=='>=' and not row['edge']>=rule['threshold']:
        return 0
    if rule['relation']=='<=' and not row['edge']<=rule['threshold']:
        return 0
    if rule['flat_cap'] is not None and not row['flat']<=rule['flat_cap']:
        return 0
    if rule['state'] is not None and row['state']!=rule['state']:
        return 0
    return rule['side']


def compare(actual,expected,label: str) -> None:
    """Проверяет равенство actual/expected с числовым допуском; label включается в ошибку."""
    if expected is None:
        assert actual is None,(label,actual,expected)
    elif isinstance(expected,float) and math.isinf(expected):
        assert actual=='∞',(label,actual,expected)
    else:
        assert actual is not None and math.isclose(actual,expected,rel_tol=1e-10,abs_tol=1e-12),(label,actual,expected)


def metrics(values: list, sides: list, mask: list) -> dict:
    """Возвращает независимо рассчитанные метрики values/sides на пригодных днях mask."""
    daily=[v for v,m in zip(values,mask) if m and v is not None]
    trades=[v for v,s,m in zip(values,sides,mask) if m and s!=0 and v is not None]
    signal_count=sum(m and s!=0 for s,m in zip(sides,mask))
    positive=[v for v in trades if v>0]
    negative=[v for v in trades if v<0]
    cumulative=peak=drawdown=0.0
    for value in daily:
        cumulative+=value
        peak=max(peak,cumulative)
        drawdown=max(drawdown,peak-cumulative)
    return {'signals':signal_count,'trades':len(trades),'missing_results':signal_count-len(trades),
            'eligible_days':sum(mask),'known_days':len(daily),
            'mean':statistics.mean(trades) if trades else None,
            'median':statistics.median(trades) if trades else None,
            'win_rate':len(positive)/len(trades) if trades else None,
            'total':sum(daily) if daily else None,'per_day':sum(daily)/len(daily) if daily else None,
            'max_drawdown':drawdown if daily else None,
            'profit_factor':sum(positive)/(-sum(negative)) if negative else (math.inf if positive else None)}


def probability_scores(rows: list, target: str, model: str, alpha: float) -> dict:
    """Независимо возвращает Brier/LogLoss/точность rows цели target для model со сглаживанием alpha."""
    labels=('Рост','Падение','Флэт') if target=='movement' else ('Бычий','Медвежий','Боковой')
    states=('Бычий','Медвежий','Боковой')
    scores,logs,hits=[],[],[]
    for row in rows:
        matrix=row['matrix_d'] if target=='movement' else row['matrix_r']
        counts=[sum(r[c] for r in matrix) for c in range(3)]
        forecast=row['p_d'] if target=='movement' else row['p_r']
        fact=row['actual_d'] if target=='movement' else row['actual_r']
        if any(v is None for v in forecast) or fact is None or sum(counts)==0 or (target=='regime' and row['state'] not in states):
            continue
        p=[(c+alpha)/(sum(counts)+3*alpha) for c in counts]
        if model=='model':
            p=forecast
        elif model=='majority':
            winner=max(range(3),key=lambda c:p[c])
            p=[int(c==winner) for c in range(3)]
        elif model=='persistence':
            p=[int(states[c]==row['state']) for c in range(3)]
        actual=labels.index(fact)
        scores.append(sum((p[c]-int(c==actual))**2 for c in range(3)))
        clipped=[max(v,1e-12) for v in p]
        logs.append(-math.log(clipped[actual]/sum(clipped)))
        winners=[c for c in range(3) if p[c]==max(p)]
        hits.append(int(len(winners)==1 and winners[0]==actual))
    return {'n':len(scores),'brier':statistics.mean(scores) if scores else None,
            'logloss':statistics.mean(logs) if logs else None,'accuracy':statistics.mean(hits) if hits else None}


def verify_summary(path: Path) -> dict:
    """Сверяет исследование JSON path с исходниками и CSV; возвращает число серий и проверенных решений."""
    path=Path(path)
    summary=json.loads(path.read_text(encoding='utf-8'))
    config=summary['config']
    scenarios=summary['scenarios']
    assert config['cost_bps'] in scenarios and max(10,2*config['cost_bps']) in scenarios
    periods=summary['periods']
    source={s['series']:s for s in summary['sources']}
    calendar=sorted({d['date'] for r in summary['results'] for d in r['decisions']})
    a,b=int(len(calendar)*0.6),int(len(calendar)*0.8)
    assert periods=={'train':calendar[:a],'validation':calendar[a:b],'test':calendar[b:]},'Неверное разделение истории'
    checked=0
    for result in summary['results']:
        series=result['series']
        original=Path(source[series]['path'])
        assert sha256(original)==source[series]['sha256'],'Исходная книга изменилась'
        rows=source_rows(original)
        assert [r['date'] for r in rows]==[d['date'] for d in result['decisions']]
        for row,decision in zip(rows,result['decisions']):
            long=signal(row,result['selection']['long'],config)
            short=signal(row,result['selection']['short'],config)
            combined=0 if long and short else long+short
            assert decision['long_side']==long and decision['short_side']==short and decision['side']==combined
            assert decision['available']==available(row,config)
            compare(decision['r_oc'],row['r'],series+' доходность')
            compare(decision['real_points'],row['real'],series+' Real')
            expected_sides={'long':long,'short':short,'combined':combined,'always_long':1,'always_short':-1,
                            'edge_sign':int(row['edge']>0)-int(row['edge']<0) if row['edge'] is not None else 0,'flat':0}
            for strategy,side in expected_sides.items():
                side=side if available(row,config) else 0
                assert decision['side_'+strategy]==side
                for cost in scenarios:
                    expected=side*row['r']-abs(side)*cost/10000 if available(row,config) and row['r'] is not None else None
                    compare(decision[f'result_{strategy}_{cost:g}bps'],expected,series+' '+strategy+' расходы')
            checked+=1
        for stored in result['summaries']:
            current=[d for d in result['decisions'] if d['period']==stored['period']]
            strategy,cost=stored['strategy'],stored['cost_bps']
            expected=metrics([d[f'result_{strategy}_{cost:g}bps'] for d in current],
                             [d['side_'+strategy] for d in current],[d['available'] for d in current])
            for key,value in expected.items():
                compare(stored[key],value,series+' '+strategy+' '+key)
        for stored in result['probabilities']['metrics']:
            current=[r for r in rows if r['date'] in periods[stored['period']]]
            expected=probability_scores(current,stored['target'],stored['model'],source[series]['metadata']['alpha'])
            for key,value in expected.items():
                compare(stored[key],value,series+' вероятность '+key)
        with (path.parent/series/'decisions.csv').open(encoding='utf-8-sig',newline='') as stream:
            csv_rows=list(csv.DictReader(stream))
        assert len(csv_rows)==len(rows),'CSV потерял строки'
        for csv_row,decision in zip(csv_rows,result['decisions']):
            assert csv_row['Дата']==decision['date']
            cell=csv_row[f'Совместные правила: результат при {config["cost_bps"]:g} б.п., доля']
            expected=decision[f'result_combined_{config["cost_bps"]:g}bps']
            compare(float(cell) if cell else None,expected,series+' CSV результат')
    return {'series':len(summary['results']),'decisions':checked,'cost_bps':config['cost_bps']}


def main(argv=None) -> int:
    """Проверяет файл --summary из argv; возвращает 0 при успехе или 1 при ошибке сверки."""
    parser=argparse.ArgumentParser(description='Независимая проверка исследования Real.')
    parser.add_argument('--summary',type=Path,default=Path(__file__).resolve().parents[1]/'outputs/real_research/summary.json',
                        help='Путь JSON готового исследования')
    args=parser.parse_args(argv)
    try:
        result=verify_summary(args.summary)
        print(f'Сверка пройдена: {result["series"]} серии, {result["decisions"]} решений; '
              f'расходы {result["cost_bps"]:g} б.п. на круг. Исходники не изменены.')
        return 0
    except (AssertionError,OSError,ValueError) as error:
        print(f'Ошибка независимой сверки: {error}')
        return 1


if __name__=='__main__':
    raise SystemExit(main())
