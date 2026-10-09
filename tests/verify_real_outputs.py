"""Независимая проверка XLSX, созданных на рабочих БД RTS и MIX.

Логика: openpyxl повторно читает результаты, SQLite открывает исходные БД
только для чтения. Проверяются календарь, десять колонок, суммы процентов,
Real=CLOSE−OPEN в пунктах, отсутствие частичных троек и текущий контракт.
Последняя строка дополнительно
сверяется с независимыми формулами по исходным ценам без импорта funds_model.

Запуск из папки проекта после создания обеих таблиц:
    python tests/verify_real_outputs.py
    python tests/verify_real_outputs.py --project C:/Users/Alkor/VSCode/pj24_funds_2 --db-dir C:/data_quote/day2
"""

import argparse
import math
import sqlite3
import statistics
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook


def verify(series: str, project: Path, db_dir: Path) -> None:
    """Сверяет XLSX series в project с БД db_dir, включая все Real; печатает итог."""
    path = project / "outputs" / f"{series}_probabilities.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        main = workbook["Прогноз"]
        assert main.max_column == 10
        assert tuple(next(main.values)) == ("Дата", "SECID", "LSTTRADE", "Рост", "Падение",
                                           "Флэт", "Бычий", "Медвежий", "Боковой", "Real")
        values = list(main.values)[1:]
        usable_d = usable_r = 0
        for row in values:
            if row[2] is not None:
                assert row[2].date() > row[0].date(), row
            for start in (3, 6):
                triple = row[start:start + 3]
                assert all(v is None for v in triple) or all(isinstance(v, (float, int)) for v in triple), row
                if triple[0] is not None:
                    assert all(0 <= v <= 1 for v in triple), row
                    assert abs(sum(triple) - 1) <= 1e-12, row
                    assert sum(round(v * 1000) for v in triple) == 1000, row
                    if start == 3:
                        usable_d += 1
                    else:
                        usable_r += 1
        audit_rows = list(workbook["Диагностика"].values)
        latest = dict(zip(audit_rows[0], audit_rows[-1]))
        settings = dict(list(workbook["Параметры"].values)[1:])
        length, k, q = settings["L"], settings["k"], settings["q"]
        target = values[-1][0].date().isoformat()
        expiry = values[-1][2].date().isoformat()
        database = (db_dir / f"{series}_day2.db").resolve()
        with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as connection:
            dates = [r[0] for r in connection.execute("SELECT DISTINCT TRADEDATE FROM Day ORDER BY TRADEDATE")]
            assert [row[0].date().isoformat() for row in values] == dates
            original = {(r[0], r[1]): r[2:] for r in connection.execute("SELECT TRADEDATE,SECID,OPEN,CLOSE,VOLUME FROM Day")}
            filled = set(connection.execute("SELECT TRADEDATE,SECID FROM FilledQuotes"))
            real_count = 0
            for row in values:
                key = (row[0].date().isoformat(), row[1])
                raw = original.get(key)
                available = (raw is not None and key not in filled and raw[2] > 0
                             and all(v is not None and math.isfinite(v) and v > 0 for v in raw[:2]))
                if available:
                    assert isinstance(row[9], (int, float)), row
                    assert math.isclose(row[9], raw[1] - raw[0], abs_tol=1e-9), row
                    real_count += 1
                else:
                    assert row[9] is None, row
            expected_expiry = connection.execute("SELECT min(LSTTRADE) FROM Day WHERE TRADEDATE<? AND LSTTRADE>?",
                                                  (target, target)).fetchone()[0]
            assert expiry == expected_expiry
            actual_secid = connection.execute("SELECT SECID FROM Day WHERE TRADEDATE=? AND LSTTRADE=? ORDER BY SECID",
                                               (target, expiry)).fetchone()[0]
            assert values[-1][1] == actual_secid
            past = connection.execute("SELECT TRADEDATE,OPEN,CLOSE,VOLUME FROM Day WHERE TRADEDATE<? AND LSTTRADE=? ORDER BY TRADEDATE DESC LIMIT 62",
                                      (target, expiry)).fetchall()[::-1]
            assert [r[0] for r in past] == dates[-63:-1]
            closes = [r[2] for r in past]
            cc = [math.log(b / a) for a, b in zip(closes, closes[1:])]
            signal_sigma = statistics.stdev(cc[:-1][-settings["W_max"]:])
            target_sigma = statistics.stdev(cc[-settings["W_max"]:])
            signal_g = math.log(closes[-1] / closes[-length - 1])
            boundary = k * signal_sigma * math.sqrt(length)
            state = "Бычий" if signal_g > boundary else "Медвежий" if signal_g < -boundary else "Боковой"
            assert latest["Наблюдаемый режим"] == state
            assert math.isclose(latest["sigma_CC предыдущего режима"], signal_sigma, rel_tol=1e-12)
            assert math.isclose(latest["sigma_CC целевого режима"], target_sigma, rel_tol=1e-12)
            assert math.isclose(latest["g предыдущего режима"], signal_g, abs_tol=1e-12)
            assert all(r[3] > 0 for r in past[-settings["V_max"]:])
            oc = [math.log(r[2] / r[1]) for r in past[-settings["V_max"]:]]
            epsilon = q * statistics.stdev(oc)
            assert math.isclose(latest["epsilon движения"], epsilon, rel_tol=1e-12)
            opening, close = connection.execute("SELECT OPEN,CLOSE FROM Day WHERE TRADEDATE=? AND LSTTRADE=?",
                                                  (target, expiry)).fetchone()
            y = math.log(close / opening)
            movement = "Рост" if y > epsilon else "Падение" if y < -epsilon else "Флэт"
            assert latest["Факт движения"] == movement
        for start in (3, 6):
            raw = audit_rows[-1][start:start + 3]
            assert all(0 <= value <= 1 for value in raw)
            assert math.isclose(sum(raw), 1, abs_tol=1e-12)
        print(f"{series}: {len(values)} дат; пригодных троек D={usable_d}, P1={usable_r}; "
              f"Real проверен для {real_count} баров; последний прогноз {target}, {actual_secid}, {expiry}; независимые формулы совпадают.")
        print("Последние отображённые проценты:", [round(v * 100, 1) for v in values[-1][3:9]])
        print("Последний Real, пунктов:", values[-1][9])
    finally:
        workbook.close()


def main() -> None:
    """Читает параметры путей и проверяет обе серии; возвращаемого значения нет."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--db-dir", type=Path, default=Path("C:/data_quote/day2"))
    args = parser.parse_args()
    for series in ("RTS", "MIX"):
        verify(series, args.project, args.db_dir)


if __name__ == "__main__":
    main()
