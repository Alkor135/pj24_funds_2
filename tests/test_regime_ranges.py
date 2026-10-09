"""Проверки подбора диапазонов и непрерывной симуляции по всей истории.

Логика: ручные примеры проверяют границы через 5%, комиссию на круг,
конфликты, пропуски, просадку от начального нуля и годовые результаты.
Независимый перебор проверяет максимум всех пар диапазонов. Геометрия столбцов
SVG проверяет годовой, а не накопленный PnL, включая убытки и нулевые годы. CLI запускается
на искусственных XLSX/SQLite; проверяется автономный HTML и сохранность входов.
Временные файлы размещаются внутри tests и автоматически удаляются.

Примеры из папки проекта:
    python -m unittest discover -s tests -p test_regime_ranges.py -v
    python -m unittest discover -s tests -v
"""

import hashlib
import importlib
import importlib.util
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from html.parser import HTMLParser
from pathlib import Path
from contextlib import closing
from xml.etree import ElementTree

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from test_real_regime_trades import trade_inputs

PROJECT = Path(__file__).resolve().parents[1]
TEST_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT))


class ReportParser(HTMLParser):
    """Извлекает встроенные данные и внешние зависимости из готового HTML."""

    def __init__(self):
        """Создаёт накопители JSON, SVG и внешних ресурсов; параметров нет."""
        super().__init__()
        self.in_data = False
        self.fragments = []
        self.external = []
        self.charts = 0
        self.yearly_charts = 0

    def handle_starttag(self, tag, attrs):
        """Обрабатывает тег tag с attrs, фиксируя данные и внешние ресурсы."""
        attributes = dict(attrs)
        if tag == "script" and attributes.get("id") == "report-data":
            self.in_data = True
        if tag == "svg" and "pnl-chart" in attributes.get("class", ""):
            self.charts += 1
        if tag == "svg" and "yearly-pnl-chart" in attributes.get("class", ""):
            self.yearly_charts += 1
        for name in ("src", "href"):
            value = attributes.get(name, "")
            if value and not value.startswith(("#", "data:")):
                self.external.append(value)

    def handle_endtag(self, tag):
        """Завершает накопление JSON при закрытии script; результат отсутствует."""
        if tag == "script":
            self.in_data = False

    def handle_data(self, data):
        """Добавляет текст data к встроенному JSON; результат отсутствует."""
        if self.in_data:
            self.fragments.append(data)


def independent_maximum(data, commission_bps):
    """Возвращает максимальный P&L независимым прямым перебором данных data.

    commission_bps — расходы на круг. Не использует производственные маски,
    кандидатов или формулы оптимизатора; каждая дневная сделка считается отдельно.
    """
    ranges = [None] + [(low / 100, high / 100)
                       for low in range(0, 100, 5) for high in range(low + 5, 101, 5)]
    rows = list(data[["Бычий", "Медвежий", "Real", "OPEN"]].itertuples(index=False, name=None))
    best = 0.0
    for bull in ranges:
        for bear in ranges:
            total = 0.0
            for p_bull, p_bear, real, opening in rows:
                buy = bull is not None and bull[0] <= p_bull and (p_bull < bull[1] or bull[1] == p_bull == 1)
                sell = bear is not None and bear[0] <= p_bear and (p_bear < bear[1] or bear[1] == p_bear == 1)
                if buy != sell and pd.notna(real):
                    total += (real if buy else -real) - opening * commission_bps / 10000
            best = max(best, total)
    return best


def configure_fixture(path, database, series, first_real=None):
    """Задаёт series в тестовых path/database и при необходимости первый first_real.

    Изменяет только искусственные файлы внутри tests; результат отсутствует.
    """
    book = load_workbook(path)
    book["Параметры"]["B2"] = series
    if first_real is not None:
        book["Прогноз"]["F2"] = first_real
    book.save(path)
    book.close()
    with closing(sqlite3.connect(database)) as connection:
        with connection:
            connection.execute("UPDATE Metadata SET VALUE=? WHERE KEY='ASSETCODE'", (series,))
            if first_real is not None:
                connection.execute("UPDATE Day SET CLOSE=OPEN+? WHERE TRADEDATE='2026-01-01'", (first_real,))


class RegimeRangesTest(unittest.TestCase):
    """Проверяет финансовые результаты и выбор диапазонов на известных примерах."""

    def module(self):
        """Возвращает модуль расчётов, проверяя наличие реализации."""
        self.assertIsNotNone(importlib.util.find_spec("source.regime_ranges"),
                             "Расчёт диапазонов ещё не реализован.")
        return importlib.import_module("source.regime_ranges")

    def sample(self):
        """Возвращает восемь дней двух лет с границами, конфликтом и пропуском Real."""
        return pd.DataFrame({
            "Дата": pd.date_range("2019-12-30", periods=8),
            "Бычий": [.05, .149, .15, .1, 1, .05, np.nan, .05],
            "Медвежий": [.1, .1, .1, .75, 0, .1, .75, .1],
            "Real": [100, -20, 999, -50, 40, 0, -30, np.nan],
            "OPEN": [10000] * 7 + [np.nan],
        })

    def test_full_history_curve_boundaries_costs_and_missing_facts(self):
        """Ловит неверные границы, знак продажи, двойную комиссию и потерю дней."""
        module = self.module()
        data = self.sample()
        before = data.copy(deep=True)
        daily = module.simulate_ranges(data, module.ProbabilityRange(5, 15),
                                       module.ProbabilityRange(70, 80), 3)
        self.assertEqual(daily["executed"].sum(), 4)
        self.assertEqual(daily["side"].tolist(), [1, 1, 0, 0, 0, 1, -1, 1])
        self.assertEqual(daily["conflict"].sum(), 1)
        self.assertEqual(daily["missing_fact"].sum(), 1)
        self.assertEqual(daily["cumulative_net"].tolist(), [97, 74, 74, 74, 74, 71, 98, 98])
        self.assertEqual(daily["cumulative_long"].iloc[-1], 71)
        self.assertEqual(daily["cumulative_short"].iloc[-1], 27)
        pd.testing.assert_frame_equal(data, before)
        metrics = module.trade_metrics(daily)
        self.assertEqual(metrics["trades"], 4)
        self.assertEqual(metrics["gross_profit"], 110)
        self.assertEqual(metrics["commission"], 12)
        self.assertEqual(metrics["net_profit"], 98)
        self.assertEqual(metrics["max_drawdown"], 26)
        self.assertEqual(metrics["win_rate"], .5)
        self.assertAlmostEqual(metrics["profit_factor"], 124 / 26)
        yearly = module.yearly_metrics(daily)
        combined = yearly[yearly["direction"] == "combined"]
        self.assertEqual(combined["year"].tolist(), [2019, 2020])
        self.assertEqual(combined["net_profit"].tolist(), [74, 24])
        self.assertEqual(combined["trades"].tolist(), [2, 2])

    def test_hundred_percent_and_drawdown_from_zero(self):
        """Ловит исключение вероятности 100% и игнорирование первого убытка."""
        module = self.module()
        data = self.sample().iloc[[4]].copy()
        data["Real"] = -40
        daily = module.simulate_ranges(data, module.ProbabilityRange(95, 100), None, 3)
        metrics = module.trade_metrics(daily)
        self.assertEqual(metrics["trades"], 1)
        self.assertEqual(metrics["net_profit"], -43)
        self.assertEqual(metrics["max_drawdown"], 43)
        self.assertEqual(metrics["profit_factor"], 0)

    def test_exhaustive_joint_maximum_matches_independent_trading(self):
        """Ловит неполный перебор или независимый выбор сторон, теряющий общий максимум."""
        module = self.module()
        generator = np.random.default_rng(24)
        data = pd.DataFrame({
            "Дата": pd.date_range("2020-01-01", periods=18),
            "Бычий": generator.integers(0, 11, 18) / 20,
            "Медвежий": generator.integers(0, 11, 18) / 20,
            "Real": generator.integers(-100, 101, 18).astype(float),
            "OPEN": generator.integers(1000, 20001, 18),
        })
        data.loc[0, "Real"] = np.nan
        data.loc[1, "Бычий"] = np.nan
        expected = independent_maximum(data, 3)
        selection = module.optimize_ranges(data, 3)
        self.assertEqual(selection["pairs_evaluated"], 44521)
        self.assertAlmostEqual(selection["net_profit"], expected, places=7)
        daily = module.simulate_ranges(data, selection["bull_range"], selection["bear_range"], 3)
        self.assertAlmostEqual(module.trade_metrics(daily)["net_profit"], expected, places=7)

    def test_no_profitable_trade_can_disable_both_directions(self):
        """Ловит принудительную убыточную торговлю при отсутствии прибыльной пары."""
        module = self.module()
        data = self.sample().iloc[[0]].copy()
        data["Real"] = 0
        selection = module.optimize_ranges(data, 3)
        self.assertIsNone(selection["bull_range"])
        self.assertIsNone(selection["bear_range"])
        self.assertEqual(selection["net_profit"], 0)
        daily = module.simulate_ranges(data, None, None, 3)
        metrics = module.trade_metrics(daily)
        self.assertEqual(metrics["trades"], 0)
        self.assertIsNone(metrics["win_rate"])
        self.assertIsNone(metrics["profit_factor"])

    def test_invalid_inputs_and_range_are_rejected(self):
        """Ловит молчаливое принятие плохих дат, вероятностей, расходов или цен."""
        module = self.module()
        for lower, upper in ((1, 10), (10, 10), (-5, 20), (5, 105)):
            with self.subTest(lower=lower, upper=upper), self.assertRaises(ValueError):
                module.ProbabilityRange(lower, upper)
        for column, value in (("Бычий", 1.1), ("Real", np.inf), ("OPEN", 0), ("Дата", None)):
            data = self.sample()
            data.loc[0, column] = value
            with self.subTest(column=column), self.assertRaises(ValueError):
                module.optimize_ranges(data, 3)
        duplicate = self.sample()
        duplicate.loc[1, "Дата"] = duplicate.loc[0, "Дата"]
        with self.assertRaises(ValueError):
            module.optimize_ranges(duplicate, 3)
        for commission in (-1, np.inf):
            with self.assertRaises(ValueError):
                module.optimize_ranges(self.sample(), commission)

    def test_cli_offline_report_has_independent_instruments_and_full_calendar(self):
        """Ловит смешение инструментов, внешние ресурсы HTML и неполную историю CLI."""
        self.module()
        script = PROJECT / "research_real_regime_ranges.py"
        self.assertTrue(script.is_file(), "Точка запуска ещё не реализована.")
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            root = Path(directory)
            mix_dir, rts_dir = root / "mix", root / "rts"
            mix_dir.mkdir()
            rts_dir.mkdir()
            mix, mix_db = trade_inputs(mix_dir)
            rts, rts_db = trade_inputs(rts_dir)
            configure_fixture(mix, mix_db, "MIX")
            configure_fixture(rts, rts_db, "RTS", first_real=-100)
            inputs = (mix, mix_db, rts, rts_db)
            before = [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs]
            output = root / "report.html"
            result = subprocess.run(
                [sys.executable, str(script), "--mix", str(mix), "--rts", str(rts),
                 "--output", str(output)], cwd=root, capture_output=True, text=True,
                encoding="utf-8", env=dict(os.environ, PYTHONIOENCODING="utf-8"), timeout=40)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            parser = ReportParser()
            parser.feed(output.read_text(encoding="utf-8"))
            self.assertEqual(parser.external, [])
            self.assertEqual(parser.charts, 8)
            self.assertEqual(parser.yearly_charts, 2)
            payload = json.loads("".join(parser.fragments))
            self.assertEqual([item["series"] for item in payload["instruments"]], ["MIX", "RTS"])
            for item, expected in zip(payload["instruments"], (172.5, 216.5)):
                self.assertEqual(len(item["daily"]), 3)
                self.assertEqual(item["pairs_evaluated"], 44521)
                self.assertAlmostEqual(item["metrics"]["combined"]["net_profit"], expected)
                self.assertEqual(len(item["yearly"]), 3)
            self.assertNotEqual(payload["instruments"][0]["bull_range"], payload["instruments"][1]["bull_range"])
            self.assertEqual(before, [hashlib.sha256(path.read_bytes()).hexdigest() for path in inputs])

    def test_yearly_bars_show_annual_profit_losses_and_zero_year(self):
        """Ловит накопительный PnL вместо годового, перепутанные стороны и знак убытка."""
        module = importlib.import_module("source.regime_ranges_report")
        self.assertTrue(hasattr(module, "yearly_svg_chart"), "Годовой столбчатый график ещё не реализован.")
        expected = {(2019, "combined"): 74, (2019, "long"): 74, (2019, "short"): 0,
                    (2020, "combined"): 24, (2020, "long"): -3, (2020, "short"): 27,
                    (2021, "combined"): 0, (2021, "long"): 0, (2021, "short"): 0}
        rows = [{"year": year, "direction": side, "net_profit": value}
                for (year, side), value in reversed(list(expected.items()))]
        svg = ElementTree.fromstring(module.yearly_svg_chart({"series": "MIX", "yearly": rows}))
        namespace = {"svg": "http://www.w3.org/2000/svg"}
        heights = {}
        for (year, side), value in expected.items():
            group = svg.find(f".//svg:g[@id='annual-MIX-{year}-{side}']", namespace)
            self.assertIsNotNone(group, f"Нет столбца {year}/{side}.")
            path = group.find("svg:path", namespace)
            coordinates = [float(number) for number in re.findall(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", path.attrib["d"])]
            heights[year, side] = coordinates[1] - coordinates[5]
        scale = heights[2019, "combined"] / 74
        self.assertGreater(scale, 0)
        for key, value in expected.items():
            self.assertAlmostEqual(heights[key] / scale, value, places=5)

    def test_cli_does_not_overwrite_database_from_workbook_metadata(self):
        """Ловит перезапись фактической SQLite из книги при совпадении её пути с HTML."""
        self.module()
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            root = Path(directory)
            path, database = trade_inputs(root)
            renamed = root / "source.html"
            database.rename(renamed)
            book = load_workbook(path)
            book["Параметры"]["B3"] = str(renamed)
            book.save(path)
            book.close()
            before = hashlib.sha256(renamed.read_bytes()).hexdigest()
            result = subprocess.run(
                [sys.executable, str(PROJECT / "research_real_regime_ranges.py"),
                 "--only", "RTS", "--rts", str(path), "--output", str(renamed)],
                cwd=root, capture_output=True, text=True, encoding="utf-8",
                env=dict(os.environ, PYTHONIOENCODING="utf-8"), timeout=40)
            self.assertEqual(result.returncode, 2, result.stdout + result.stderr)
            self.assertEqual(before, hashlib.sha256(renamed.read_bytes()).hexdigest())


if __name__ == "__main__":
    unittest.main()
