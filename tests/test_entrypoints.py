"""Проверки запуска корневых скриптов из другой рабочей папки.

Логика: копирует программные файлы в временный проект внутри tests, запускает
настоящие точки входа отдельным процессом и проверяет стандартные пути XLSX
и исследования. Использует искусственные SQLite/XLSX; рабочие данные не меняет.
Временный проект автоматически удаляется. Также проверяет --help всех запусков.

Примеры из папки проекта:
    python -m unittest discover -s tests -p test_entrypoints.py -v
    python -m unittest discover -s tests -v
"""

import json
import os
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import closing
from pathlib import Path

from openpyxl import load_workbook

from test_funds_model import database_fixture, history
from test_real_research import workbook_fixture


PROJECT = Path(__file__).resolve().parents[1]
TEST_ROOT = Path(__file__).resolve().parent


def copy_program(destination: Path) -> Path:
    """Копирует скрипты и пакет source в destination; возвращает корень копии."""
    destination.mkdir()
    for script in PROJECT.glob("*.py"):
        shutil.copy2(script, destination / script.name)
    shutil.copytree(PROJECT / "source", destination / "source",
                    ignore=shutil.ignore_patterns("__pycache__"))
    return destination


def launch(script: Path, cwd: Path, *arguments: str) -> subprocess.CompletedProcess:
    """Запускает script с arguments из cwd; возвращает код выхода и текст потоков."""
    environment = dict(os.environ, PYTHONIOENCODING="utf-8")
    return subprocess.run([sys.executable, str(script), *arguments], cwd=cwd,
                          capture_output=True, text=True, encoding="utf-8",
                          env=environment, timeout=60)


class EntrypointsTest(unittest.TestCase):
    """Проверяет импорты и пути результатов реальных точек входа после переноса."""

    def test_help_from_another_directory(self):
        """Проверяет --help трёх скриптов из другой папки; результата нет."""
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            for script in ("forecast_mix.py", "forecast_rts.py", "research_real.py"):
                with self.subTest(script=script):
                    result = launch(PROJECT / script, Path(directory), "--help")
                    self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                    self.assertIn("--help", result.stdout)

    def test_forecast_default_outputs_from_another_directory(self):
        """Проверяет создание обеих XLSX в outputs корня копии; результата нет."""
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            root = Path(directory)
            project = copy_program(root / "project")
            for series in ("MIX", "RTS"):
                with self.subTest(series=series):
                    database = root / f"{series}.db"
                    database_fixture(database, history())
                    with closing(sqlite3.connect(database)) as connection:
                        with connection:
                            connection.execute("UPDATE Metadata SET VALUE=? WHERE KEY='ASSETCODE'",
                                               (series,))
                    result = launch(project / f"forecast_{series.lower()}.py", root,
                                    "--db", str(database))
                    self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                    output = project / "outputs" / f"{series}_probabilities.xlsx"
                    self.assertTrue(output.is_file(), result.stdout)
                    book = load_workbook(output, read_only=True, data_only=True)
                    try:
                        self.assertEqual(book["Прогноз"].max_row, 81)
                        self.assertEqual(book["Прогноз"].max_column, 10)
                    finally:
                        book.close()

    def test_research_default_paths_from_another_directory(self):
        """Проверяет стандартные входы и артефакты исследования; результата нет."""
        with tempfile.TemporaryDirectory(dir=TEST_ROOT) as directory:
            root = Path(directory)
            project = copy_program(root / "project")
            outputs = project / "outputs"
            outputs.mkdir()
            for series in ("MIX", "RTS"):
                workbook_fixture(outputs / f"{series}_probabilities.xlsx", series)
            report = root / "report.md"
            result = launch(project / "research_real.py", root,
                            "--bootstrap-repeats", "2", "--report", str(report))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            summary = outputs / "real_research" / "summary.json"
            self.assertTrue(summary.is_file(), result.stdout)
            data = json.loads(summary.read_text(encoding="utf-8"))
            self.assertEqual([item["series"] for item in data["results"]], ["MIX", "RTS"])
            self.assertTrue(report.is_file())


if __name__ == "__main__":
    unittest.main()
