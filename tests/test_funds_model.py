"""Проверки методики дневных вероятностей и создаваемых XLSX.

Логика: искусственная история проверяет формулы, отсутствие будущих данных,
переименование контрактов, экспирацию, разрывы и синтетические цены.
Временные SQLite/XLSX проверяют чтение, колонку Real в пунктах и сохранение
без изменения исходных БД. Общие модули импортируются из пакета source.

Запуск из папки проекта:
    python -m unittest discover -s tests -v
"""

import math
import sqlite3
import statistics
import tempfile
import unittest
import io
from contextlib import redirect_stdout
from contextlib import closing
from datetime import date, timedelta
from pathlib import Path


def history(count=80, expiry=date(2030, 3, 21), secid="OLD", scale=100.0):
    """Создаёт count баров контракта expiry/secid с ценовым масштабом scale.

    Возвращает список Quote с изменяющимися дневными доходностями.
    """
    from source.funds_model import Quote

    result = []
    close = scale
    for i in range(count):
        opening = close * math.exp(0.002 * math.cos(i))
        close *= math.exp(0.008 * math.sin(i * 0.7) + 0.001)
        result.append(Quote(date(2020, 1, 1) + timedelta(days=i), secid,
                            expiry, opening, close, 100, False))
    return result


def database_fixture(path: Path, quotes) -> None:
    """Создаёт в path временную БД RTS из quotes; результата нет."""
    with closing(sqlite3.connect(path)) as connection:
        connection.execute("CREATE TABLE Day(TRADEDATE,SECID,LSTTRADE,OPEN,CLOSE,VOLUME)")
        connection.executemany("INSERT INTO Day VALUES(?,?,?,?,?,?)",
            [(q.day.isoformat(), q.secid, q.expiry.isoformat(), q.open, q.close, q.volume)
             for q in quotes])
        connection.execute("CREATE TABLE Metadata(KEY,VALUE)")
        connection.execute("INSERT INTO Metadata VALUES('ASSETCODE','RTS')")
        connection.commit()


class ModelTest(unittest.TestCase):
    """Проверяет поведение модели на искусственных данных без параметров."""

    def test_regime_formula_and_shift(self):
        """Сверяет первое окно L=20 и исключение текущей доходности из sigma."""
        from source.funds_model import Config, regime_feature

        closes = [q.close for q in history(21)]
        returns = [math.log(b / a) for a, b in zip(closes, closes[1:])]
        feature = regime_feature(closes, Config())
        self.assertEqual(feature.window, 19)
        self.assertAlmostEqual(feature.sigma, statistics.stdev(returns[:-1]), 14)
        self.assertAlmostEqual(feature.g, math.log(closes[-1] / closes[0]), 14)
        self.assertIsNone(regime_feature(closes[:-1], Config()).state)
        self.assertIsNone(regime_feature([100.0] * 62, Config()).state)

    def test_short_windows_and_gap(self):
        """Проверяет полный прогрев, разрыв и минимальную выборку волатильности."""
        from source.funds_model import Config, regime_feature, volatility

        self.assertEqual(regime_feature([q.close for q in history(62)], Config()).window, 60)
        self.assertIsNone(regime_feature([100.0] * 30 + [None] + [101.0] * 20, Config()).state)
        self.assertEqual(volatility([0.1, 0.2, None, 0.3], 60).window, 1)
        self.assertIsNone(volatility([0.1, 0.1], 60).sigma)

    def test_movement_boundaries(self):
        """Проверяет строгие внешние границы и принадлежность равенства Flat."""
        from source.funds_model import classify_movement

        self.assertEqual(classify_movement(0.25, 0.25), 2)
        self.assertEqual(classify_movement(-0.25, 0.25), 2)
        self.assertEqual(classify_movement(0.251, 0.25), 0)
        self.assertEqual(classify_movement(-0.251, 0.25), 1)

    def test_smoothing_and_fallback(self):
        """Сверяет сглаживание alpha=0.5, резерв и отсутствие истории."""
        from source.funds_model import estimate

        matrix = [[52, 28, 20], [0, 0, 0], [0, 0, 0]]
        probabilities, source, count = estimate(matrix, 0, 0.5)
        self.assertAlmostEqual(probabilities[0], 52.5 / 101.5)
        self.assertEqual((source, count), ("CONDITIONAL_STATE", 100))
        self.assertEqual(estimate(matrix, 1, 0.5),
                         (probabilities, "FALLBACK_UNCONDITIONAL", 100))
        self.assertEqual(estimate(matrix, None, 0.5),
                         (probabilities, "FALLBACK_UNCONDITIONAL", 100))
        self.assertIsNone(estimate([[0] * 3 for _ in range(3)], 0, 0.5)[0])

    def test_future_prices_do_not_change_forecasts(self):
        """Изменение целевого и будущих баров не меняет предшествующий прогноз."""
        from dataclasses import replace
        from source.funds_model import build_forecasts

        quotes = history()
        reference = build_forecasts(quotes)
        changed = [replace(q, open=q.open * 3, close=q.close * 2)
                   if i >= 40 else q for i, q in enumerate(quotes)]
        alternative = build_forecasts(changed)
        prefix = build_forecasts(quotes[:41])
        for i in range(41):
            self.assertEqual(reference[i].movement, alternative[i].movement)
            self.assertEqual(reference[i].regime, alternative[i].regime)
            self.assertEqual(reference[i].movement, prefix[i].movement)
            self.assertEqual(reference[i].regime, prefix[i].regime)

    def test_secid_alias_preserves_expiry_history(self):
        """Переименование SECID сохраняет вероятности и выводит алиас из БД."""
        from dataclasses import replace
        from source.funds_model import build_forecasts

        quotes = history()
        renamed = [replace(q, secid="NEW") if i >= 40 else q
                   for i, q in enumerate(quotes)]
        left, right = build_forecasts(quotes), build_forecasts(renamed)
        self.assertEqual(right[40].secid, "NEW")
        self.assertEqual([r.movement for r in left], [r.movement for r in right])
        self.assertEqual([r.regime for r in left], [r.regime for r in right])

    def test_rollover_and_one_training_example(self):
        """Экспирация переключает контракт без склейки цен и удвоения обучения."""
        from source.funds_model import build_forecasts

        old = history(expiry=date(2020, 2, 10))[:41]
        new = history(expiry=date(2030, 3, 21), secid="NEW", scale=10000.0)
        result = build_forecasts(old + new)
        self.assertEqual(result[40].secid, "NEW")
        self.assertEqual(result[40].expiry, date(2030, 3, 21))
        self.assertLessEqual(result[-1].details["train_D"], 80 - 22)
        self.assertLessEqual(result[-1].details["train_R"], 80 - 22)
        self.assertAlmostEqual(result[40].details["signal_g"],
                               math.log(new[39].close / new[19].close))

    def test_missing_bar_does_not_compress_calendar(self):
        """Пропуск одного контракта сохраняет общую дату и обрывает его окно."""
        from source.funds_model import build_forecasts

        main = history()
        backup = history(expiry=date(2030, 6, 20), secid="BACKUP")
        result = build_forecasts(main[:35] + main[36:] + backup)
        self.assertIsNone(result[36].details["signal_state"])
        self.assertIsNone(result[36].regime)
        self.assertIsNone(result[36].movement)

    def test_filled_quote_excludes_movement_only(self):
        """Перенесённый бар исключается из движения, но сохраняет режимный исход."""
        from dataclasses import replace
        from source.funds_model import build_forecasts

        quotes = history()
        reference = build_forecasts(quotes)
        modified = quotes.copy()
        modified[35] = replace(quotes[35], filled=True, volume=0)
        changed = build_forecasts(modified)
        self.assertEqual(changed[36].details["train_D"], reference[36].details["train_D"] - 1)
        self.assertEqual(changed[36].details["train_R"], reference[36].details["train_R"])

    def test_real_points_are_signed_and_available_without_training(self):
        """Real сохраняет CLOSE−OPEN, включая первый бар и ноль, без обучения."""
        from source.funds_model import Quote, build_forecasts

        quotes = [Quote(date(2020, 1, i + 1), "X", date(2030, 3, 21), opening,
                        close, volume, filled)
                  for i, (opening, close, volume, filled) in enumerate([
                      (100.25, 104.75, 1, False), (105, 100, 1, False),
                      (101, 101, 1, False), (100, 105, 0, True),
                      (100, 105, 0, False)])]
        rows = build_forecasts(quotes)
        self.assertEqual([r.details.get("real_points") for r in rows], [4.5, -5, 0, None, None])
        self.assertIsNone(rows[0].movement)

    def test_alias_duplicates_and_conflicts(self):
        """Одинаковые алиасы не удваивают обучение, конфликт цен вызывает ошибку."""
        from dataclasses import replace
        from source.funds_model import build_forecasts

        quotes = history()
        duplicate = replace(quotes[35], secid="SECOND")
        self.assertEqual(build_forecasts(quotes)[-1].movement,
                         build_forecasts(quotes + [duplicate])[-1].movement)
        with self.assertRaises(ValueError):
            build_forecasts(quotes + [replace(duplicate, close=1.0)])

    def test_future_date_requires_fresh_previous_session(self):
        """Будущая дата требует совпадения ожидаемой прошлой сессии с концом БД."""
        from source.funds_model import build_forecasts

        quotes = history()
        target = quotes[-1].day + timedelta(days=1)
        with self.assertRaises(ValueError):
            build_forecasts(quotes, forecast_date=target)
        with self.assertRaises(ValueError):
            build_forecasts(quotes, forecast_date=target,
                            previous_session=quotes[-1].day + timedelta(days=1))
        result = build_forecasts(quotes, forecast_date=target,
                                 previous_session=quotes[-1].day)
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0].day, target)
        self.assertIsNotNone(result[0].movement)

    def test_read_database_marks_filled_quotes(self):
        """Чтение SQLite распознаёт FilledQuotes и проверяет принадлежность серии."""
        from source.funds_model import read_database

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            path = Path(directory) / "test.db"
            with closing(sqlite3.connect(path)) as connection:
                connection.execute("CREATE TABLE Day(TRADEDATE,SECID,LSTTRADE,OPEN,CLOSE,VOLUME)")
                connection.execute("INSERT INTO Day VALUES('2020-01-01','X','2030-03-21',100,101,0)")
                connection.execute("CREATE TABLE FilledQuotes(TRADEDATE,SECID)")
                connection.execute("INSERT INTO FilledQuotes VALUES('2020-01-01','X')")
                connection.execute("CREATE TABLE Metadata(KEY,VALUE)")
                connection.execute("INSERT INTO Metadata VALUES('ASSETCODE','RTS')")
                connection.commit()
            self.assertTrue(read_database(path, "RTS")[0].filled)
            with self.assertRaises(ValueError):
                read_database(path, "MIX")


class WorkbookTest(unittest.TestCase):
    """Проверяет округление и сохранённую структуру XLSX без параметров."""

    def test_rounding_preserves_displayed_sum(self):
        """Каждая отображаемая тройка составляет 100.0% при фиксированном порядке."""
        from source.funds_xlsx import rounded_probabilities

        self.assertEqual(rounded_probabilities((1 / 3,) * 3), (0.334, 0.333, 0.333))
        self.assertAlmostEqual(sum(rounded_probabilities((52.5 / 101.5, 28.5 / 101.5, 20.5 / 101.5))), 1)

    def test_xlsx_has_requested_columns_and_numeric_values(self):
        """XLSX содержит десять колонок, Флэт, проценты, Real и диагностику."""
        from openpyxl import load_workbook
        from source.funds_model import build_forecasts
        from source.funds_xlsx import write_workbook

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            path = Path(directory) / "test.xlsx"
            write_workbook(path, build_forecasts(history()), {"Серия": "RTS"})
            workbook = load_workbook(path)
            sheet = workbook["Прогноз"]
            self.assertEqual(tuple(c.value for c in sheet[1]),
                             ("Дата", "SECID", "LSTTRADE", "Рост", "Падение", "Флэт",
                              "Бычий", "Медвежий", "Боковой", "Real"))
            self.assertEqual(sheet.max_column, 10)
            self.assertEqual(sheet.max_row, 81)
            self.assertEqual(sheet["D81"].number_format, "0.0%")
            self.assertEqual(sheet["A81"].number_format, "yyyy-mm-dd")
            self.assertIsInstance(sheet["D81"].value, float)
            self.assertAlmostEqual(sum(sheet.cell(81, c).value for c in range(4, 7)), 1)
            source = history()[-1]
            self.assertAlmostEqual(sheet["J81"].value, source.close - source.open, 12)
            self.assertEqual(sheet["J81"].number_format, "#,##0.00")
            self.assertIn("Диагностика", workbook.sheetnames)
            self.assertIn("Параметры", workbook.sheetnames)
            workbook.close()
            with self.assertRaises(FileExistsError):
                write_workbook(path, [], {})

    def test_cli_filters_preserve_training(self):
        """Фильтры CLI сокращают вывод, сохраняя обучение по более ранним датам."""
        from openpyxl import load_workbook
        from source.funds_model import build_forecasts, run_cli
        from source.funds_xlsx import rounded_probabilities

        quotes = history()
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            database = Path(directory) / "quotes.db"
            output = Path(directory) / "filtered.xlsx"
            database_fixture(database, quotes)
            with redirect_stdout(io.StringIO()):
                status = run_cli("RTS", ["--db", str(database), "--output", str(output),
                    "--start-date", quotes[40].day.isoformat(), "--end-date", quotes[45].day.isoformat()])
            self.assertEqual(status, 0)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["Прогноз"].max_row, 7)
                actual = tuple(workbook["Прогноз"].cell(2, column).value for column in range(4, 7))
                self.assertEqual(actual, rounded_probabilities(build_forecasts(quotes)[40].movement))
            finally:
                workbook.close()

    def test_cli_future_and_stale_data(self):
        """CLI создаёт одну будущую строку и отклоняет устаревшую БД без XLSX."""
        from openpyxl import load_workbook
        from source.funds_model import run_cli

        quotes = history()
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as directory:
            database = Path(directory) / "quotes.db"
            output = Path(directory) / "future.xlsx"
            database_fixture(database, quotes)
            target = quotes[-1].day + timedelta(days=1)
            arguments = ["--db", str(database), "--output", str(output),
                         "--forecast-date", target.isoformat()]
            with redirect_stdout(io.StringIO()):
                self.assertEqual(run_cli("RTS", arguments), 2)
                self.assertFalse(output.exists())
                self.assertEqual(run_cli("RTS", arguments + ["--previous-session", quotes[-1].day.isoformat()]), 0)
            workbook = load_workbook(output)
            try:
                self.assertEqual(workbook["Прогноз"].max_row, 2)
                self.assertEqual(workbook["Прогноз"]["A2"].value.date(), target)
                self.assertIsNone(workbook["Прогноз"]["J2"].value)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
